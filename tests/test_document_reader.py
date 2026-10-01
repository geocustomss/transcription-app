import io
import unittest

from docx import Document
from pypdf import PdfWriter

from src.document_reader import DocumentInputError, extract_document_text


class DocumentReaderTests(unittest.TestCase):
    def test_extracts_utf8_text_and_markdown(self):
        self.assertEqual(extract_document_text("notes.txt", b"\xef\xbb\xbfHello, world."), "Hello, world.")
        self.assertEqual(extract_document_text("notes.md", "# Meeting\n\nNotes".encode()), "# Meeting\n\nNotes")

    def test_extracts_docx_paragraphs_and_tables_in_order(self):
        document = Document()
        document.add_paragraph("Agenda")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Speaker"
        table.cell(0, 1).text = "Action"
        buffer = io.BytesIO()
        document.save(buffer)

        self.assertEqual(
            extract_document_text("minutes.docx", buffer.getvalue()),
            "Agenda\nSpeaker\tAction",
        )

    def test_reports_scanned_or_empty_text_pdf(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)

        with self.assertRaisesRegex(DocumentInputError, "Scanned PDFs need OCR"):
            extract_document_text("scan.pdf", buffer.getvalue())

    def test_rejects_empty_unsupported_and_oversized_documents(self):
        with self.assertRaisesRegex(DocumentInputError, "empty"):
            extract_document_text("notes.txt", b"")
        with self.assertRaisesRegex(DocumentInputError, "Unsupported document type"):
            extract_document_text("notes.rtf", b"text")
        with self.assertRaisesRegex(DocumentInputError, "limit"):
            extract_document_text("notes.txt", b"long", max_bytes=3)

    def test_rejects_documents_over_the_text_limit(self):
        with self.assertRaisesRegex(DocumentInputError, "character limit"):
            extract_document_text("notes.txt", b"four", max_characters=3)


if __name__ == "__main__":
    unittest.main()